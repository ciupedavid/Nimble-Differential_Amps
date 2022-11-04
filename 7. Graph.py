from openpyxl import load_workbook
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series
)

import json
# This script creates a scatter graph
with open(r'paths.json') as d:
    paths = json.load(d)['Unzip'][0]

workbook_path = (paths['transfer_function'])
workbook = load_workbook(workbook_path)
sheet = workbook.active

# Range of the axis form the excel
x_nimble = Reference(sheet, min_col=2, min_row=2, max_row=1000)
y_nimble = Reference(sheet, min_col=1, min_row=2, max_row=1000)
x_ltspice = Reference(sheet, min_col=4, min_row=2, max_row=1000)
y_ltspice = Reference(sheet, min_col=3, min_row=2, max_row=1000)
x_datasheet = Reference(sheet, min_col=6, min_row=2, max_row=1000)
y_datasheet = Reference(sheet, min_col=5, min_row=2, max_row=1000)

# Series name
series_voltage = Series(x_nimble, y_nimble,
                     title_from_data=False, title="Nimble")
series_freq = Series(x_ltspice, y_ltspice,
                     title_from_data=False, title="LTspice")
series_mag = Series(x_datasheet, y_datasheet,
                     title_from_data=False, title="Datasheet")

# Chart type
chart = ScatterChart()
chart.series.append(series_voltage)
chart.series.append(series_freq)
chart.series.append(series_mag)
chart.x_axis.scaling.logBase = 10

chart.x_axis.scaling.min = 100000
chart.y_axis.scaling.min = -15
chart.x_axis.scaling.max = 10000000000
chart.y_axis.scaling.max = 10
chart.x_axis.tickLblPos = "low"
#chart.x_axis.tickLblSkip = 3

chart.title = None
chart.x_axis.title = 'Frequency (Hz)'
chart.y_axis.title = 'Magnitude (dB)'
chart.legend.position = 'r'

sheet.add_chart(chart, 'J1')
workbook.save(workbook_path)